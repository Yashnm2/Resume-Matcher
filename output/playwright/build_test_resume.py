from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("jane-doe-test-resume.docx")


def set_font(run, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.375)
    paragraph.paragraph_format.first_line_indent = Inches(-0.188)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    set_font(paragraph.add_run(text), size=10.5)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, before, after in [
    ("Heading 1", 16, 10, 5),
    ("Heading 2", 13, 8, 3),
]:
    style = doc.styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string("1F4D78")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

name = doc.add_paragraph()
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.paragraph_format.space_after = Pt(1)
set_font(name.add_run("JANE DOE"), size=22, bold=True, color="0B2545")

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(3)
set_font(title.add_run("Senior Backend Engineer"), size=12, bold=True, color="2E74B5")

contact = doc.add_paragraph()
contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
contact.paragraph_format.space_after = Pt(8)
set_font(
    contact.add_run(
        "jane@example.com | +1-555-0100 | San Francisco, CA | janedoe.dev | github.com/janedoe"
    ),
    size=9.5,
    color="444444",
)

doc.add_heading("Professional Summary", level=1)
doc.add_paragraph(
    "Backend engineer with 6 years of experience building scalable Python APIs, "
    "microservices, and reliable cloud platforms."
)

doc.add_heading("Experience", level=1)
role = doc.add_paragraph()
role.paragraph_format.space_after = Pt(2)
set_font(role.add_run("Senior Backend Engineer — Acme Corp"), size=11, bold=True)
set_font(role.add_run("  |  Jan 2021 - Present"), size=10, color="555555")
add_bullet(doc, "Built FastAPI services handling more than 50,000 requests per day.")
add_bullet(doc, "Led migration from a monolith to a resilient microservices architecture.")
add_bullet(doc, "Mentored three junior developers on testing and backend best practices.")

role = doc.add_paragraph()
role.paragraph_format.space_before = Pt(4)
role.paragraph_format.space_after = Pt(2)
set_font(role.add_run("Software Engineer — StartupCo"), size=11, bold=True)
set_font(role.add_run("  |  Jun 2018 - Dec 2020"), size=10, color="555555")
add_bullet(doc, "Developed a payment system processing $2M in monthly transactions.")
add_bullet(doc, "Raised automated test coverage from 40% to 85%.")

doc.add_heading("Education", level=1)
education = doc.add_paragraph()
education.paragraph_format.space_after = Pt(3)
set_font(education.add_run("MIT — B.S. Computer Science"), size=11, bold=True)
set_font(education.add_run("  |  2014 - 2018"), size=10, color="555555")
doc.add_paragraph("Graduated with honors; Dean's List.")

doc.add_heading("Projects", level=1)
project = doc.add_paragraph()
project.paragraph_format.space_after = Pt(2)
set_font(project.add_run("OpenAPI Generator — Creator & Maintainer"), size=11, bold=True)
add_bullet(doc, "Built a CLI that generates API clients from OpenAPI specifications.")
add_bullet(doc, "Earned 500+ GitHub stars and adoption by more than 30 companies.")

doc.add_heading("Skills", level=1)
skills = doc.add_paragraph()
skills.paragraph_format.space_after = Pt(0)
set_font(skills.add_run("Python, FastAPI, Docker, AWS, PostgreSQL, Redis, REST APIs, Microservices"), size=10.5)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run("Synthetic test résumé — no real personal data"), size=8, color="777777")

doc.core_properties.title = "Jane Doe Synthetic Test Resume"
doc.core_properties.author = "Resume Matcher automated test"
doc.save(OUT)
print(OUT)
